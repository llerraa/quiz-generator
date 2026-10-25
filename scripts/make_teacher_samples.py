"""
Делает набор учебных файлов «как от преподавателей» — с разными способами
отметки правильного ответа. Нужен, чтобы проверить режим «Тест для курса»
на настоящей модели, пока реальных файлов ещё нет.

Запуск из корня проекта:
    python scripts/make_teacher_samples.py

Файлы появятся в data/sample/teacher_files/. Все вопросы выдуманные
(основы Python), персональных данных нет — файлы можно хранить в git.

Что ожидать (для сверки результата в приложении):
    01_жирный.docx            10 из 10 — верный вариант выделен жирным
    02_строка_ответ.docx      10 из 10 — после вариантов «Ответ: Б»
    03_звёздочка.txt          10 из 10 — звёздочка перед верным вариантом
    04_ключ_в_конце.docx      10 из 10 — ключ «1-б, 2-а…» отдельным списком
    05_таблица.docx           10 из 10 — вопрос, варианты и ответ в таблице
    06_цвет_подчёркивание.docx 9 из 9 — зелёный цвет / подчёркивание / маркер
    07_жирный.pdf             10 из 10 — то же, что 01, но в PDF
    08_сложный.docx           3 из 3 в тест; 4 пропущено: нет ответа, два ответа,
                              открытый вопрос, 7 вариантов
    09_aiken_30.txt           25 в тест (первые), ещё 5 не вошли
    10_кодировка_1251.txt     10 из 10 — «Ответ: Б», файл в Windows-1251
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "sample" / "teacher_files"

# (вопрос, [варианты], индекс верного)
POOL = [
    ("Какой символ начинает комментарий в Python?", ["//", "#", "--", "/*"], 1),
    ("Какой тип данных у значения 3.14?", ["int", "str", "float", "bool"], 2),
    ("Как получить длину списка a?", ["a.size()", "len(a)", "length(a)", "a.count"], 1),
    ("Какое ключевое слово объявляет функцию?", ["func", "define", "def", "function"], 2),
    ("Что выведет print(2 ** 3)?", ["6", "8", "9", "5"], 1),
    ("Какая структура данных хранит пары «ключ — значение»?", ["list", "tuple", "set", "dict"], 3),
    ("Каким оператором проверяют равенство?", ["=", "==", "=>", "!="], 1),
    ("Как называется неизменяемая последовательность?", ["list", "dict", "tuple", "set"], 2),
    ("Какой цикл выполняется, пока условие истинно?", ["for", "while", "do", "loop"], 1),
    ("Что делает метод append списка?", ["удаляет элемент", "добавляет элемент в конец", "сортирует список", "копирует список"], 1),
    ("Какая функция читает строку с клавиатуры?", ["scan()", "read()", "input()", "get()"], 2),
    ("Какой модуль нужен для работы с регулярными выражениями?", ["regex", "re", "match", "rx"], 1),
]
LETTERS = "абвг"


def _doc_to(path, doc):
    doc.save(str(path))


def make_bold(path):
    doc = Document()
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        doc.add_paragraph(f"{n}. {q}")
        for i, o in enumerate(opts):
            run = doc.add_paragraph().add_run(f"{LETTERS[i]}) {o}")
            run.bold = i == ok
    _doc_to(path, doc)


def make_answer_line(path):
    doc = Document()
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        doc.add_paragraph(f"Вопрос {n}. {q}")
        for i, o in enumerate(opts):
            doc.add_paragraph(f"{LETTERS[i]}) {o}")
        doc.add_paragraph(f"Ответ: {LETTERS[ok].upper()}")
    _doc_to(path, doc)


def make_star(path):
    lines = []
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        lines.append(f"{n}. {q}")
        for i, o in enumerate(opts):
            lines.append(f"{'*' if i == ok else ' '}{LETTERS[i]}) {o}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def make_key_at_end(path):
    doc = Document()
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        doc.add_paragraph(f"{n}. {q}")
        for i, o in enumerate(opts):
            doc.add_paragraph(f"{LETTERS[i]}) {o}")
    doc.add_paragraph("")
    doc.add_paragraph("Ключ к тесту")
    doc.add_paragraph(", ".join(f"{n}-{LETTERS[ok]}" for n, (_, _, ok) in enumerate(POOL[:10], 1)))
    _doc_to(path, doc)


def make_table(path):
    doc = Document()
    table = doc.add_table(rows=1, cols=7)
    for cell, title in zip(table.rows[0].cells, ["№", "Вопрос", "А", "Б", "В", "Г", "Верный"]):
        cell.text = title
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(n), q
        for i, o in enumerate(opts):
            cells[2 + i].text = o
        cells[6].text = LETTERS[ok].upper()
    _doc_to(path, doc)


def make_colour_underline(path):
    doc = Document()
    for n, (q, opts, ok) in enumerate(POOL[:9], 1):
        doc.add_paragraph(f"{n}. {q}")
        for i, o in enumerate(opts):
            run = doc.add_paragraph().add_run(f"{LETTERS[i]}) {o}")
            if i == ok:
                style = n % 3
                if style == 0:
                    run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
                elif style == 1:
                    run.underline = True
                else:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _doc_to(path, doc)


def make_bold_pdf(path):
    fonts = next(
        (d for d in (ROOT / "assets", ROOT / "tests" / "assets", ROOT / "assets" / "fonts")
         if (d / "DejaVuSans.ttf").exists() and (d / "DejaVuSans-Bold.ttf").exists()),
        None,
    )
    if fonts is None:
        raise SystemExit("Нужны шрифты DejaVuSans.ttf и DejaVuSans-Bold.ttf в папке assets/ (или tests/assets/).")
    pdfmetrics.registerFont(TTFont("DejaVuSans", fonts / "DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", fonts / "DejaVuSans-Bold.ttf"))
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        if y < 150:
            c.showPage()
            y = 800
        c.setFont("DejaVuSans", 11)
        c.drawString(50, y, f"{n}. {q}")
        y -= 18
        for i, o in enumerate(opts):
            c.setFont("DejaVuSans-Bold" if i == ok else "DejaVuSans", 11)
            c.drawString(70, y, f"{LETTERS[i]}) {o}")
            y -= 16
        y -= 10
    c.save()


def make_hard(path):
    doc = Document()
    # 1. нормальный
    q, opts, ok = POOL[0]
    doc.add_paragraph(f"1. {q}")
    for i, o in enumerate(opts):
        doc.add_paragraph(f"{LETTERS[i]}) {o}")
    doc.add_paragraph(f"Ответ: {LETTERS[ok].upper()}")
    # 2. без отмеченного ответа
    q, opts, _ = POOL[1]
    doc.add_paragraph(f"2. {q}")
    for i, o in enumerate(opts):
        doc.add_paragraph(f"{LETTERS[i]}) {o}")
    # 3. два верных ответа
    doc.add_paragraph("3. Какие из перечисленных — числовые типы данных?")
    for i, (o, good) in enumerate([("int", True), ("str", False), ("float", True), ("list", False)]):
        doc.add_paragraph().add_run(f"{LETTERS[i]}) {o}").bold = good
    # 4. открытый вопрос
    doc.add_paragraph("4. Объясните своими словами, чем список отличается от кортежа.")
    # 5. нормальный, ответ жирным
    q, opts, ok = POOL[3]
    doc.add_paragraph(f"5. {q}")
    for i, o in enumerate(opts):
        doc.add_paragraph().add_run(f"{LETTERS[i]}) {o}").bold = i == ok
    # 6. семь вариантов
    doc.add_paragraph("6. Выберите самый популярный язык для анализа данных:")
    for i, o in enumerate(["Python", "Go", "Rust", "C", "Java", "Ruby", "PHP"]):
        doc.add_paragraph(f"{i + 1}) {o}")
    doc.add_paragraph("Ответ: 1")
    # 7. нормальный, звёздочка
    q, opts, ok = POOL[4]
    doc.add_paragraph(f"7. {q}")
    for i, o in enumerate(opts):
        doc.add_paragraph(f"{'* ' if i == ok else ''}{LETTERS[i]}) {o}")
    _doc_to(path, doc)


def make_aiken_30(path):
    blocks = []
    for n in range(30):
        q, opts, ok = POOL[n % len(POOL)]
        lines = [f"{q} (вариант {n + 1})"] + [f"{'ABCD'[i]}. {o}" for i, o in enumerate(opts)] + [f"ANSWER: {'ABCD'[ok]}"]
        blocks.append("\n".join(lines))
    path.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")


def make_cp1251(path):
    lines = []
    for n, (q, opts, ok) in enumerate(POOL[:10], 1):
        lines.append(f"{n}. {q}")
        lines.extend(f"{LETTERS[i]}) {o}" for i, o in enumerate(opts))
        lines.append(f"Ответ: {LETTERS[ok].upper()}")
        lines.append("")
    path.write_bytes("\n".join(lines).encode("cp1251", errors="replace"))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    makers = {
        "01_жирный.docx": make_bold,
        "02_строка_ответ.docx": make_answer_line,
        "03_звёздочка.txt": make_star,
        "04_ключ_в_конце.docx": make_key_at_end,
        "05_таблица.docx": make_table,
        "06_цвет_подчёркивание.docx": make_colour_underline,
        "07_жирный.pdf": make_bold_pdf,
        "08_сложный.docx": make_hard,
        "09_aiken_30.txt": make_aiken_30,
        "10_кодировка_1251.txt": make_cp1251,
    }
    for name, maker in makers.items():
        maker(OUT / name)
        print("создан", OUT / name)


if __name__ == "__main__":
    main()
