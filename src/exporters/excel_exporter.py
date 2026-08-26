"""
Экспорт сгенерированных вопросов в Excel (.xlsx) через openpyxl.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models import (
    EssayQuestion,
    MatchingQuestion,
    MultipleChoiceQuestion,
    Question,
    ShortAnswerQuestion,
    TrueFalseQuestion,
)

_HEADER_FILL = PatternFill(start_color="FC4D0F", end_color="FC4D0F", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(wrap_text=True, vertical="top")

_TYPE_LABELS_RU = {
    "multiple_choice": "Множественный выбор",
    "true_false": "Верно/неверно",
    "matching": "На соответствие",
    "essay": "Эссе",
    "short_answer": "Короткий ответ",
}

_COLUMNS = [
    "№", "Тип вопроса", "Вопрос / задание", "Варианты ответов",
    "Правильный ответ", "Объяснение", "Страницы",
]


def _question_row(index: int, q: Question) -> list:
    pages = ", ".join(str(p) for p in q.source_pages) if q.source_pages else ""

    if isinstance(q, MultipleChoiceQuestion):
        options_text = "\n".join(f"{chr(65 + i)}) {opt}" for i, opt in enumerate(q.options))
        correct_text = f"{chr(65 + q.correct_index)}) {q.options[q.correct_index]}"
        return [
            index, _TYPE_LABELS_RU["multiple_choice"], q.question,
            options_text, correct_text, q.explanation or "", pages,
        ]

    if isinstance(q, TrueFalseQuestion):
        return [
            index, _TYPE_LABELS_RU["true_false"], q.statement,
            "", "ВЕРНО" if q.is_true else "НЕВЕРНО", "", pages,
        ]

    if isinstance(q, MatchingQuestion):
        pairs_text = "\n".join(f"{p.left} → {p.right}" for p in q.pairs)
        return [
            index, _TYPE_LABELS_RU["matching"], q.instruction,
            pairs_text, pairs_text, "", pages,
        ]

    if isinstance(q, EssayQuestion):
        # У эссе нет единственного "правильного ответа" — вместо этого
        # критерии оценки, по которым преподаватель проверяет вручную.
        return [
            index, _TYPE_LABELS_RU["essay"], q.question,
            "", "", q.grading_criteria, pages,
        ]

    if isinstance(q, ShortAnswerQuestion):
        return [
            index, _TYPE_LABELS_RU["short_answer"], q.question,
            "", q.sample_answer, "", pages,
        ]

    return [index, "Неизвестный тип вопроса", "", "", "", "", pages]


def questions_to_excel_bytes(questions: list[Question]) -> bytes:
    """Собирает Excel-файл с вопросами и возвращает его как байты."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Тест"

    ws.append(_COLUMNS)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for i, q in enumerate(questions, start=1):
        ws.append(_question_row(i, q))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP

    for i, width in enumerate([5, 20, 40, 35, 25, 35, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()