"""Тесты для src/exporters/excel_exporter.py."""
from io import BytesIO

from openpyxl import load_workbook

from src.exporters.excel_exporter import questions_to_excel_bytes
from src.models import MatchingPair, MatchingQuestion, MultipleChoiceQuestion, TrueFalseQuestion


def test_questions_to_excel_produces_readable_file():
    questions = [
        MultipleChoiceQuestion(
            question="Что такое ReLU?",
            options=["Функция активации", "Тип слоя"],
            correct_index=0,
            explanation="Определение из лекции",
            source_pages=[1, 2],
        ),
    ]
    data = questions_to_excel_bytes(questions)

    # Открываем обратно тем же openpyxl — если файл битый, load_workbook упадёт
    wb = load_workbook(BytesIO(data))
    ws = wb.active

    assert ws["A1"].value == "№"
    assert ws["C2"].value == "Что такое ReLU?"
    assert ws["D2"].value == "A) Функция активации\nB) Тип слоя"
    assert ws["E2"].value == "A) Функция активации"


def test_questions_to_excel_true_false_shows_verno_neverno():
    questions = [TrueFalseQuestion(statement="Утверждение", is_true=False, source_pages=[3])]
    wb = load_workbook(BytesIO(questions_to_excel_bytes(questions)))
    ws = wb.active

    assert ws["C2"].value == "Утверждение"
    assert ws["E2"].value == "НЕВЕРНО"
    assert ws["G2"].value == "3"


def test_questions_to_excel_matching_lists_pairs():
    questions = [
        MatchingQuestion(
            instruction="Сопоставьте",
            pairs=[MatchingPair(left="A", right="1"), MatchingPair(left="B", right="2")],
        )
    ]
    wb = load_workbook(BytesIO(questions_to_excel_bytes(questions)))
    ws = wb.active

    assert ws["D2"].value == "A → 1\nB → 2"


def test_questions_to_excel_handles_multiple_questions_and_row_count():
    questions = [
        MultipleChoiceQuestion(question="Q1?", options=["A", "B"], correct_index=0),
        TrueFalseQuestion(statement="Q2", is_true=True),
    ]
    wb = load_workbook(BytesIO(questions_to_excel_bytes(questions)))
    ws = wb.active

    assert ws.max_row == 3  # заголовок + 2 вопроса
    assert ws["A2"].value == 1
    assert ws["A3"].value == 2


def test_questions_to_excel_empty_list_still_produces_valid_file():
    wb = load_workbook(BytesIO(questions_to_excel_bytes([])))
    assert wb.active.max_row == 1  # только заголовок