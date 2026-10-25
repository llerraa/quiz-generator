"""
Тесты: оформление (жирный, подчёркивание, маркер, цвет) превращается в метки,
а таблицы читаются в порядке документа. Нужно, чтобы приложение видело
правильный ответ, отмеченный не словами, а оформлением.
"""
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from src.parsing.docx_parser import parse_docx
from src.parsing.pdf_parser import parse_pdf

_ROOT = Path(__file__).resolve().parent.parent
# Шрифт DejaVu нужен только для тестов с PDF (в стандартных шрифтах reportlab нет кириллицы).
# Ищем в нескольких привычных местах — положите оба .ttf в любое из них.
# Имена шрифтов в reportlab: свои, уникальные. reportlab молча игнорирует повторную регистрацию
# занятого имени, поэтому с общими именами («DejaVuSans-Bold») тест ломался бы, если другой
# модуль успел зарегистрировать это имя раньше на другой файл.
_FONT_REGULAR = "ImesEmphasisTestSans"
_FONT_BOLD = "ImesEmphasisTestSans-Bold"
_FONT_DIR_CANDIDATES = [
    Path(__file__).parent / "assets",
    _ROOT / "assets",
    _ROOT / "assets" / "fonts",
]


@pytest.fixture(scope="module")
def pdf_fonts():
    for folder in _FONT_DIR_CANDIDATES:
        regular, bold = folder / "DejaVuSans.ttf", folder / "DejaVuSans-Bold.ttf"
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont(_FONT_REGULAR, regular))
            pdfmetrics.registerFont(TTFont(_FONT_BOLD, bold))
            return
    pytest.skip(
        "Нет шрифтов DejaVuSans.ttf и DejaVuSans-Bold.ttf — положите их в папку assets/ "
        "(или tests/assets/), и тесты с PDF заработают."
    )


def _save(doc, tmp_path, name="q.docx"):
    path = tmp_path / name
    doc.save(str(path))
    return path


# ---------- DOCX ----------

def test_bold_underline_highlight_color_become_marks(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("а) обычный")
    doc.add_paragraph().add_run("б) жирный").bold = True
    doc.add_paragraph().add_run("в) подчёркнутый").underline = True
    doc.add_paragraph().add_run("г) маркер").font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph().add_run("д) красный").font.color.rgb = RGBColor(0xFF, 0, 0)

    text = parse_docx(_save(doc, tmp_path))
    lines = text.splitlines()
    assert lines[0] == "а) обычный"
    assert lines[1] == "**б) жирный**"
    assert lines[2] == "__в) подчёркнутый__"
    assert lines[3] == "==г) маркер=="
    assert lines[4] == "{~д) красный~}"


def test_black_font_color_is_not_emphasis(tmp_path):
    doc = Document()
    doc.add_paragraph().add_run("обычный чёрный").font.color.rgb = RGBColor(0, 0, 0)
    assert parse_docx(_save(doc, tmp_path)) == "обычный чёрный"


def test_only_part_of_paragraph_is_marked_and_runs_are_merged(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Ответ: ")
    r1 = p.add_run("Пи")
    r1.bold = True
    r2 = p.add_run("тон")  # Word часто режет слово на куски — метки не должны дробиться
    r2.bold = True
    text = parse_docx(_save(doc, tmp_path))
    assert text == "Ответ: **Питон**"


def test_spaces_stay_outside_marks(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("слово")
    r = p.add_run(" жирное ")
    r.bold = True
    p.add_run("дальше")
    assert parse_docx(_save(doc, tmp_path)) == "слово **жирное** дальше"


def test_mark_emphasis_can_be_turned_off(tmp_path):
    doc = Document()
    doc.add_paragraph().add_run("жирный").bold = True
    assert parse_docx(_save(doc, tmp_path), mark_emphasis=False) == "жирный"


def test_bold_paragraph_style_counts_as_bold(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    bold_style = doc.styles.add_style("ЖирныйСтиль", WD_STYLE_TYPE.PARAGRAPH)
    bold_style.font.bold = True
    doc.add_paragraph("через стиль", style="ЖирныйСтиль")
    doc.add_paragraph("без стиля")
    # жирность задана стилем абзаца, а не самим текстом — метка всё равно нужна
    assert parse_docx(_save(doc, tmp_path)).splitlines() == ["**через стиль**", "без стиля"]


def test_table_is_read_in_document_order_and_as_rows(tmp_path):
    doc = Document()
    doc.add_paragraph("До таблицы")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "1"
    table.rows[0].cells[1].text = "Вопрос?"
    table.rows[0].cells[2].text = "Ответ: Б"
    table.rows[1].cells[0].text = "2"
    table.rows[1].cells[1].text = "Ещё вопрос?"
    table.rows[1].cells[2].paragraphs[0].add_run("А").bold = True
    doc.add_paragraph("После таблицы")

    lines = parse_docx(_save(doc, tmp_path)).splitlines()
    assert lines == [
        "До таблицы",
        "1 | Вопрос? | Ответ: Б",
        "2 | Ещё вопрос? | **А**",
        "После таблицы",
    ]


def test_merged_cells_are_not_duplicated(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    merged = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    merged.text = "Общая ячейка"
    assert parse_docx(_save(doc, tmp_path)) == "Общая ячейка"


def test_hyperlink_paragraph_falls_back_to_plain_text(tmp_path):
    from docx.oxml import OxmlElement

    doc = Document()
    p = doc.add_paragraph()
    p.add_run("до ").bold = True
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "ссылка"
    run.append(t)
    link.append(run)
    p._p.append(link)
    # Текст гиперссылки не входит в обычные фрагменты абзаца. Метки в таком абзаце
    # не ставим, но и слова терять нельзя.
    assert parse_docx(_save(doc, tmp_path)) == "до ссылка"


# ---------- PDF ----------

def _make_pdf(path, lines):
    """lines: список (текст, bold, rgb|None)."""
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 780
    for text, bold, rgb in lines:
        c.setFont(_FONT_BOLD if bold else _FONT_REGULAR, 12)
        c.setFillColorRGB(*(rgb or (0, 0, 0)))
        c.drawString(50, y, text)
        y -= 22
    c.showPage()
    c.save()
    return path


def test_pdf_bold_and_coloured_words_are_marked(tmp_path, pdf_fonts):
    pdf = _make_pdf(tmp_path / "q.pdf", [
        ("1. Что такое переменная?", False, None),
        ("а) тип цикла", False, None),
        ("б) имя для данных", True, None),
        ("в) красный вариант", False, (1, 0, 0)),
    ])
    text = parse_pdf(pdf, mark_emphasis=True).full_text
    assert "**б) имя для данных**" in text
    assert "{~в) красный вариант~}" in text
    assert "а) тип цикла" in text and "**а) тип цикла" not in text


def test_pdf_without_emphasis_is_unchanged(tmp_path, pdf_fonts):
    pdf = _make_pdf(tmp_path / "plain.pdf", [
        ("Обычная строка один", False, None),
        ("Обычная строка два", False, None),
    ])
    marked = parse_pdf(pdf, mark_emphasis=True).full_text
    plain = parse_pdf(pdf).full_text
    assert marked == plain
    assert "**" not in marked


def test_pdf_marks_are_off_by_default(tmp_path, pdf_fonts):
    pdf = _make_pdf(tmp_path / "b.pdf", [("жирная строка", True, None)])
    assert "**" not in parse_pdf(pdf).full_text