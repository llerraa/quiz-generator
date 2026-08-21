"""Тесты для src/parsing/docx_parser.py."""
import pytest
from docx import Document

from src.parsing.docx_parser import parse_docx


@pytest.fixture
def sample_docx(tmp_path):
    """Создаёт тестовый docx с вопросом в обычном абзаце и таблицей."""
    doc = Document()
    doc.add_paragraph("1. Что такое риск?")
    doc.add_paragraph("а) Вероятность потерь (верно)")
    doc.add_paragraph("б) Гарантия прибыли")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Термин"
    table.rows[0].cells[1].text = "Определение"

    path = tmp_path / "test_questions.docx"
    doc.save(str(path))
    return path


def test_parse_docx_extracts_paragraph_text(sample_docx):
    text = parse_docx(sample_docx)
    assert "Что такое риск?" in text
    assert "Вероятность потерь (верно)" in text


def test_parse_docx_extracts_table_text(sample_docx):
    text = parse_docx(sample_docx)
    assert "Термин" in text
    assert "Определение" in text


def test_parse_docx_raises_on_missing_file():
    with pytest.raises(FileNotFoundError):
        parse_docx("data/sample/does_not_exist.docx")


def test_parse_docx_raises_value_error_on_corrupt_file(tmp_path):
    fake_docx = tmp_path / "not_really.docx"
    fake_docx.write_text("это не настоящий docx-файл")
    with pytest.raises(ValueError):
        parse_docx(fake_docx)