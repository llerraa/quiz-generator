import io
from pathlib import Path

import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from src.parsing.file_loader import decode_text_file, load_question_file


def test_utf8_and_utf8_with_bom():
    assert decode_text_file("Вопрос".encode("utf-8")) == "Вопрос"
    assert decode_text_file("Вопрос".encode("utf-8-sig")) == "Вопрос"


def test_windows_1251_text_is_decoded_correctly():
    raw = "Какой язык интерпретируемый?".encode("cp1251")
    assert decode_text_file(raw) == "Какой язык интерпретируемый?"


def test_txt_goes_through_decoder():
    loaded = load_question_file("A. да".encode("cp1251"), ".TXT")
    assert loaded.text == "A. да"
    assert loaded.warnings == []


def test_docx_bytes_keep_emphasis_marks():
    doc = Document()
    doc.add_paragraph().add_run("правильный").bold = True
    buf = io.BytesIO()
    doc.save(buf)
    loaded = load_question_file(buf.getvalue(), ".docx")
    assert loaded.text == "**правильный**"


def test_pdf_without_text_layer_gives_warning(tmp_path):
    pdf = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    c.showPage()  # пустая страница = «скан» без текста
    c.save()
    loaded = load_question_file(pdf.read_bytes(), ".pdf")
    assert any("нет текстового слоя" in w for w in loaded.warnings)


def test_unsupported_format_raises():
    with pytest.raises(ValueError, match="не поддерживается"):
        load_question_file(b"x", ".doc")


def test_broken_docx_raises_value_error():
    with pytest.raises(ValueError):
        load_question_file(b"not a docx", ".docx")
